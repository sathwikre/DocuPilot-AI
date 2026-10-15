"""History: each run keeps its uploaded file and outputs, can be reopened, updated, deleted and cleared."""
import io
import json
from dataclasses import replace

import pytest
from openpyxl import load_workbook

import config.settings as cs
import utils.history as h
from core.agent_orchestrator import rebuild


@pytest.fixture()
def hist_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(h, "SETTINGS", replace(cs.SETTINGS, reports_dir=tmp_path, history_enabled=True, history_files=True))
    return tmp_path


def test_run_stores_source_and_outputs(hist_dir, invoice, invoice_result):
    res = replace(invoice_result)
    run_id = h.record_run(res, invoice[0])
    assert res.run_id == run_id
    entry = h.load_history()[-1]
    files = h.run_files(entry)
    assert set(files) == {"source", "workbook", "validation", "result"}
    assert files["source"].read_bytes() == invoice[0]
    assert files["workbook"].read_bytes() == res.excel_bytes
    assert len(json.loads(files["validation"].read_text())) == len(res.issues)
    assert entry["critical"] == res.metrics()["critical"] and entry["options"] == sorted(res.options)


def test_reopened_run_is_the_same_result(hist_dir, invoice, invoice_result):
    h.record_run(replace(invoice_result), invoice[0])
    back = h.load_result(h.load_history()[-1])
    assert back is not None and back.run_id == h.load_history()[-1]["run_id"]
    assert back.excel_bytes == invoice_result.excel_bytes
    assert [t.values for t in back.tables] == [t.values for t in invoice_result.tables]


def test_tampered_result_is_not_loaded(hist_dir, invoice, invoice_result):
    h.record_run(replace(invoice_result), invoice[0])
    entry = h.load_history()[-1]
    path = h.run_files(entry)["result"]
    path.write_bytes(path.read_bytes() + b"x")
    assert h.load_result(entry) is None


def test_apply_updates_the_stored_workbook(hist_dir, invoice, invoice_result):
    res = replace(invoice_result)
    h.record_run(res, invoice[0])
    new = rebuild(res, res.options - {"flag_cells"})
    assert new.run_id == res.run_id
    h.update_run(new)
    entry = h.load_history()[-1]
    assert entry["options"] == sorted(new.options) and "updated" in entry
    ws = load_workbook(io.BytesIO(h.run_files(entry)["workbook"].read_bytes()))["Extracted Data"]
    assert not any(c.comment for row in ws.iter_rows() for c in row)
    assert h.load_result(entry).options == new.options


def test_delete_and_clear(hist_dir, invoice, invoice_result):
    a = h.record_run(replace(invoice_result), invoice[0])
    b = h.record_run(replace(invoice_result), invoice[0])
    assert a != b, "two runs of the same file get distinct ids"
    h.delete_run(a)
    assert [e["run_id"] for e in h.load_history()] == [b]
    assert not (hist_dir / "history" / a).exists() and (hist_dir / "history" / b).exists()
    h.clear_history()
    assert h.load_history() == [] and not (hist_dir / "history").exists()


def test_metrics_only_when_file_storage_is_off(tmp_path, monkeypatch, invoice, invoice_result):
    monkeypatch.setattr(h, "SETTINGS", replace(cs.SETTINGS, reports_dir=tmp_path, history_files=False))
    h.record_run(replace(invoice_result), invoice[0])
    entry = h.load_history()[-1]
    assert "files" not in entry and h.run_files(entry) == {}
    assert not (tmp_path / "history").exists()


def test_run_ids_cannot_escape_the_history_folder(hist_dir):
    with pytest.raises(ValueError):
        h.run_dir("../../etc")
    assert h.run_files({"run_id": "../x", "files": {"source": "a"}}) == {}
