"""Clean-up checklist: detection counts, each fix, instant rebuild, and the upload options (OCR mode, repeats)."""
import copy
import dataclasses
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from core import cleanup
from core.agent_orchestrator import process_document, rebuild
from core.data_normalizer import type_table
from core.models import LogicalTable, RowMeta
from core.validator import Validator, check_summary


def raw_table(header, rows, method="pdf-ruled", table_id="T1"):
    return LogicalTable(table_id=table_id, title=None, pages=[1], original_header=list(header), columns=list(header),
                        roles=[None] * len(header), rows=[list(r) for r in rows],
                        row_meta=[RowMeta(row_id=f"{table_id}-R{i + 1}", page=1, method=method) for i in range(len(rows))],
                        methods=[method])


def typed(t, selected=cleanup.DEFAULT_SELECTION):
    ref = raw_table(t.original_header, t.rows)
    type_table(ref)
    out = cleanup.clean([t], selected, {ref.table_id: ref.col_types})
    for x in out.tables:
        type_table(x, opts=cleanup.typing_options(selected))
    return out


ITEMS = [["Item", "Group", "Amount"],
         ["Alpha", "NORTH REGION", "1.234,50"],
         ["Beta", "", "2.000"],
         ["Gamma", "", "15,25"],
         ["Delta", "SOUTH REGION", "7,00"],
         ["Delta", "SOUTH REGION", "7,00"],
         ["Epsilon", "", "3,10"]]


def test_detect_counts_and_suggestions():
    t = raw_table(ITEMS[0], ITEMS[1:])
    type_table(t)
    found = cleanup.detect([t])
    assert found["duplicates"]["count"] == 1 and found["duplicates"]["suggested"]
    assert found["decimal_comma"]["suggested"] and found["decimal_comma"]["count"] >= 2
    assert set(found) == set(cleanup.OPTION_IDS)
    assert all(len(v["examples"]) <= 3 for v in found.values())


def test_duplicates_are_removed_and_recorded():
    out = typed(raw_table(ITEMS[0], ITEMS[1:]), cleanup.DEFAULT_SELECTION | {"duplicates"})
    t = out.tables[0]
    assert len(t.rows) == 5
    assert [r.row_id for r in out.removed] == ["T1-R5"]
    assert "T1-R4" in out.removed[0].reason


def test_decimal_comma_option_reads_european_numbers():
    plain = typed(raw_table(ITEMS[0], ITEMS[1:])).tables[0]
    euro = typed(raw_table(ITEMS[0], ITEMS[1:]), cleanup.DEFAULT_SELECTION | {"decimal_comma"}).tables[0]
    assert euro.values[0][2] == Decimal("1234.50")
    assert euro.values[1][2] == Decimal("2000")
    assert euro.values[2][2] == Decimal("15.25")
    assert plain.values[1][2] == Decimal("2.000"), "without the option a lone dot is a decimal point"


def test_case_and_fill_down():
    rows = [["NORTH REGION", "Pens", "10"], ["", "Paper", "20"], ["", "Ink", "30"],
            ["SOUTH REGION", "Pens", "15"], ["", "Paper", "25"], ["WEST REGION", "Ink", "5"]]
    sel = cleanup.DEFAULT_SELECTION | {"case", "fill_down"}
    t = typed(raw_table(["Region", "Item", "Qty"], rows), sel).tables[0]
    assert t.rows[0][0] == "North Region"
    assert [r[0] for r in t.rows] == ["North Region"] * 3 + ["South Region"] * 2 + ["West Region"]
    assert any("filled into the blank cell" in op for op in t.row_meta[1].ops)
    found = cleanup.detect([typed(raw_table(["Region", "Item", "Qty"], rows)).tables[0]])
    assert found["fill_down"]["count"] == 3 and found["fill_down"]["suggested"]
    assert found["case"]["suggested"] is False, "only 3 of 9 text cells are capitals"


def test_text_fixes():
    rows = [["A1", "manufac- turing  plant", "10"], ["A2", "GST  invoice", "20"], ["A3", "plain", "30"]]
    t = typed(raw_table(["Code", "Note", "Qty"], rows), cleanup.DEFAULT_SELECTION | {"hyphen_breaks"}).tables[0]
    assert t.rows[0][1] == "manufacturing plant"
    assert t.rows[1][1] == "GST invoice"


def test_drcr_sign_and_zero_blanks():
    rows = [["01/05/2026", "Opening", "100.00 Cr"], ["02/05/2026", "Fee", "20.00 Dr"],
            ["03/05/2026", "Deposit", ""], ["04/05/2026", "Fee", "5.00 Dr"], ["05/05/2026", "Fee", "1.00 Dr"]]
    t = typed(raw_table(["Date", "Detail", "Amount"], rows), cleanup.DEFAULT_SELECTION | {"drcr", "zero_blanks"}).tables[0]
    assert t.values[0][2] == Decimal("100.00") and t.values[1][2] == Decimal("-20.00")
    assert t.values[2][2] == Decimal(0)


def test_month_first_dates():
    rows = [["03/04/2026", "x", "1"], ["05/06/2026", "y", "2"], ["07/08/2026", "z", "3"]]
    dmy = typed(raw_table(["Date", "Item", "Qty"], rows)).tables[0]
    mdy = typed(raw_table(["Date", "Item", "Qty"], rows), cleanup.DEFAULT_SELECTION | {"month_first"}).tables[0]
    assert dmy.values[0][0] == date(2026, 4, 3)
    assert mdy.values[0][0] == date(2026, 3, 4)


def test_numbers_and_dates_can_stay_text():
    rows = [["01/05/2026", "1,200.00"], ["02/05/2026", "300.00"], ["03/05/2026", "40.00"]]
    t = typed(raw_table(["Date", "Amount"], rows), frozenset()).tables[0]
    assert t.col_types == ["text", "text"]
    assert t.values[0] == ["01/05/2026", "1,200.00"]


def test_blank_rows_and_empty_columns_dropped():
    rows = [["a", "", "1"], ["", "", ""], ["b", "", "2"], ["c", "", "3"]]
    out = typed(raw_table(["Name", "", "Qty"], rows))
    assert out.tables[0].original_header == ["Name", "Qty"]
    assert len(out.tables[0].rows) == 3
    assert out.removed[0].reason == "Blank row"


def test_date_order_check_flags_a_misread_date():
    rows = [[f"{d:02d}/05/2026", "Item", "1"] for d in (1, 2, 3, 4, 25, 6, 7, 8)]
    t = raw_table(["Date", "Item", "Qty"], rows)
    type_table(t)
    issues = Validator().validate([t], [], "generic")
    flagged = [i for i in issues if i.check == "date_order" and i.severity == "warning"]
    assert [i.row_id for i in flagged] == ["T1-R6"]
    summary = {c["check"]: c for c in check_summary(issues)}
    assert summary["date_order"]["checked"] == 7 and summary["date_order"]["passed"] == 6


def test_rebuild_is_instant_and_does_not_change_the_original(invoice_result):
    before = invoice_result.excel_bytes
    r = rebuild(invoice_result, invoice_result.options | {"case"})
    assert r is not invoice_result and invoice_result.excel_bytes == before
    assert r.rebuild_seconds is not None and r.rebuild_seconds < 5
    assert r.metrics()["critical"] == invoice_result.metrics()["critical"]
    assert r.options == invoice_result.options | {"case"}


def test_totals_move_to_their_own_sheet(invoice_result):
    base = copy.deepcopy(invoice_result.base_tables)
    t = base[0]
    t.rows.append(["Total"] + [""] * (len(t.original_header) - 1))
    t.row_meta.append(RowMeta(row_id="T1-R99", page=1, method=t.row_meta[0].method))
    with_total = dataclasses.replace(invoice_result, base_tables=base)
    r = rebuild(with_total, invoice_result.options | {"totals_separate"})
    wb = load_workbook(io.BytesIO(r.excel_bytes))
    assert "Totals" in wb.sheetnames
    data_ids = {wb["Extracted Data"].cell(i, 1).value for i in range(2, wb["Extracted Data"].max_row + 1)}
    totals_ids = {c.value for c in wb["Totals"]["A"]}
    assert "T1-R99" in totals_ids and "T1-R99" not in data_ids
    kept = rebuild(with_total, invoice_result.options)
    assert "Totals" not in load_workbook(io.BytesIO(kept.excel_bytes)).sheetnames


def test_rebuild_without_flags_has_no_comments(invoice_result):
    r = rebuild(invoice_result, invoice_result.options - {"flag_cells"})
    ws = load_workbook(io.BytesIO(r.excel_bytes))["Extracted Data"]
    assert not any(c.comment for row in ws.iter_rows() for c in row)


def test_rebuild_rejects_unknown_options(invoice_result):
    with pytest.raises(ValueError):
        rebuild(invoice_result, {"not_an_option"})


def test_checklist_and_options_on_result(invoice_result):
    assert invoice_result.options == cleanup.default_selection(True)
    assert set(invoice_result.checklist) == set(cleanup.OPTION_IDS)
    wb = load_workbook(io.BytesIO(invoice_result.excel_bytes))
    labels = [c.value for c in wb["Document Summary"]["A"] if c.value]
    assert "Checks" in labels and "Clean-up applied" in labels


def test_keep_repeats_option(invoice):
    r = process_document(invoice[0], "inv.pdf", remove_repeats=False)
    assert "duplicates" not in r.options


def test_ocr_mode_force_reads_digital_pages_with_ocr(invoice, ocr_ok):
    if not ocr_ok:
        pytest.skip("OCR engine not available")
    r = process_document(invoice[0], "inv.pdf", ocr_mode="force")
    assert r.ocr_pages == [1] and r.ocr_mode == "force"


@pytest.mark.parametrize("raw, clean", [
    ("**Total**", "Total"), ("**Grand Total**", "Grand Total"), ("__Due__", "Due"), ("**", ""),
    ("**Grand", "Grand"), ("Total**", "Total"), ("Total**:", "Total:"), ("Amount* (see note)", "Amount* (see note)"),
    ("2*3", "2*3"), ("plain text", "plain text"), ("**1,200.00**", "1,200.00"), ("snake_case_id", "snake_case_id"),
])
def test_strip_markup(raw, clean):
    from core.value_parser import strip_markup
    assert strip_markup(raw) == clean


def test_markdown_markers_never_reach_the_workbook(invoice_result):
    base = copy.deepcopy(invoice_result.base_tables)
    t = base[0]
    desc = next(c for c, h in enumerate(t.original_header) if "desc" in h.lower())
    qty = next(c for c, h in enumerate(t.original_header) if "qty" in h.lower() or "quantity" in h.lower())
    t.original_header[desc] = f"**{t.original_header[desc]}**"
    t.rows[0][desc] = f"**{t.rows[0][desc]}**"
    t.rows[0][qty] = f"**{t.rows[0][qty]}**"
    r = rebuild(dataclasses.replace(invoice_result, base_tables=base), invoice_result.options)
    out = r.tables[0]
    assert "**" not in "".join(out.columns) and not any("**" in c for row in out.rows for c in row)
    assert out.values[0][qty] == invoice_result.tables[0].values[0][qty], "a bold number is still a number"
    ws = load_workbook(io.BytesIO(r.excel_bytes))["Extracted Data"]
    assert not any(isinstance(c.value, str) and "**" in c.value for row in ws.iter_rows() for c in row)


def test_markdown_markers_removed_from_words():
    from core.agent_orchestrator import _strip_word_markup
    from core.models import Word
    words = [Word("**Total**", 0, 0, 1, 1), Word("**", 1, 0, 2, 1), Word("1,200.00", 2, 0, 3, 1)]
    assert [w.text for w in _strip_word_markup(words)] == ["Total", "1,200.00"]


def test_ocr_mode_off_reports_scanned_pages(statement_scanned):
    r = process_document(statement_scanned[0], "st.pdf", ocr_mode="off")
    assert any("switched off" in w for w in r.warnings)
    assert r.failed_pages and not r.ocr_pages
