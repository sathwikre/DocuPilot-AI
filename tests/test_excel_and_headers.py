"""Excel workbook correctness and header normalisation."""
import io
from datetime import datetime

from openpyxl import load_workbook

from core.header_normalizer import header_similarity, normalise_headers


def wb_of(result):
    return load_workbook(io.BytesIO(result.excel_bytes))


def test_invoice_workbook_structure(invoice_result):
    wb = wb_of(invoice_result)
    assert wb.sheetnames[:4] == ["Extracted Data", "Validation Report", "Extraction Audit", "Document Summary"]
    ws = wb["Extracted Data"]
    header = [c.value for c in ws[1]]
    assert header[0] == "Row ID" and header[-1] == "Source Page"
    assert header[1:-1] == invoice_result.tables[0].columns
    assert ws.freeze_panes == "B2"
    assert ws.tables, "data must be an Excel table with filter buttons"


def test_numeric_types_and_formats(invoice_result):
    ws = wb_of(invoice_result)["Extracted Data"]
    header = [c.value for c in ws[1]]
    lt = header.index("Line Total") + 1
    cell = ws.cell(2, lt)
    assert isinstance(cell.value, float) and cell.value == 247.80
    assert "0.00" in cell.number_format
    neg = ws.cell(15, lt)
    assert neg.value == -138.05
    disc = ws.cell(2, header.index("Disc %") + 1)
    assert disc.value is None, "legitimately empty cells stay empty"
    tax = ws.cell(2, header.index("Tax %") + 1)
    assert abs(tax.value - 0.18) < 1e-12 and "%" in tax.number_format


def test_flagged_cell_has_comment(invoice_result):
    ws = wb_of(invoice_result)["Extracted Data"]
    header = [c.value for c in ws[1]]
    bad = ws.cell(8, header.index("Line Total") + 1)  # row 7 of the data
    assert bad.comment is not None and "CRITICAL" in bad.comment.text


def test_validation_and_audit_sheets_match_results(invoice_result):
    wb = wb_of(invoice_result)
    vr = wb["Validation Report"]
    assert vr.max_row - 1 == len(invoice_result.issues)
    headers = [c.value for c in vr[1]]
    for col in ("Check", "Severity", "Status", "Table", "Row", "Expected Value", "Extracted Value", "Difference", "Explanation"):
        assert col in headers
    au = wb["Extraction Audit"]
    assert au.max_row - 1 == sum(len(t.rows) for t in invoice_result.tables)
    assert [c.value for c in au[1]][:5] == ["Source Page", "Table", "Row ID", "Row Type", "Extraction Method"]


def test_multi_table_workbook_has_sheet_per_table(report_result):
    wb = wb_of(report_result)
    assert len(wb.sheetnames) == 4 + len(report_result.tables)
    assert any(name.startswith("T1") for name in wb.sheetnames)


def test_dates_written_as_dates(report_result):
    wb = wb_of(report_result)
    ws = next(wb[n] for n in wb.sheetnames if n.startswith("T3"))
    header = [c.value for c in ws[1]]
    v = ws.cell(2, header.index("Invoice Date") + 1).value
    assert isinstance(v, datetime)


def test_formula_injection_is_neutralised():
    from core.excel_exporter import build_workbook
    from core.data_normalizer import type_table
    from core.models import DocumentInfo, LogicalTable, ProcessingResult, RowMeta
    t = LogicalTable("T1", None, [1], ["A", "B"], ["A", "B"], [None, None],
                     [["=HYPERLINK(\"http://x\")", "x"], ["y", "z"]],
                     [RowMeta("T1-R1", 1, "t"), RowMeta("T1-R2", 1, "t")], ["t"])
    type_table(t)
    res = ProcessingResult("f.pdf", "h", DocumentInfo("f.pdf", "pdf", 1, []), "generic", [], [t], [], [], [], {}, 0, [])
    ws = load_workbook(io.BytesIO(build_workbook(res)))["Extracted Data"]
    c = ws.cell(2, 2)
    assert c.data_type == "s" and c.value.startswith("=HYPERLINK")


def test_header_normalisation_rules():
    names, roles, ops = normalise_headers(["Txn Date", "Descripton", "Debit Amount", "Credit Amount", "Bal"])
    assert names == ["Transaction Date", "Description", "Debit Amount", "Credit Amount", "Balance"]
    assert roles == ["transaction_date", "description", "debit", "credit", "balance"]
    assert ops, "original headers must be recorded"
    names, roles, _ = normalise_headers(["Posting Date", "Value Date"])
    assert roles == ["posting_date", "value_date"], "distinct concepts must stay distinct"


def test_header_similarity_tolerates_ocr_noise():
    assert header_similarity(["Txn Date", "Description", "Balance"], ["Txn Date", "Descripton", "Balanse"]) > 0.85
    assert header_similarity(["SKU", "Item", "Category"], ["Vendor", "Invoice No", "Status"]) < 0.5
