"""Excel workbook generation (openpyxl).

Sheets:
  1. Extracted Data      - every table, typed values, filters, frozen header, number formats.
                           Cells with validation findings are highlighted and carry a comment.
  2. Validation Report   - one row per check result.
  3. Extraction Audit    - one row per extracted row: page, table, method, operations, warnings.
  4. Document Summary    - pages, merge decisions, header mappings, extracted fields.
  5+ T<n> sheets         - one sheet per table when the document has more than one table, each with
                           its own frozen header and auto-filter.
Only metadata that exists is written; nothing is invented.
"""
from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal
from typing import Any, Optional, Sequence

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.models import LogicalTable, ProcessingResult, ValidationIssue

HEADER_FILL = PatternFill("solid", fgColor="1E1B4B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13, color="1E1B4B")
META_FONT = Font(italic=True, size=9, color="64748B")
ISSUE_FILL = {"critical": PatternFill("solid", fgColor="FECACA"), "warning": PatternFill("solid", fgColor="FEF3C7"),
              "info": PatternFill("solid", fgColor="DBEAFE"), "passed": PatternFill("solid", fgColor="DCFCE7")}
THIN = Border(bottom=Side(style="thin", color="CBD5E1"))
_SAFE_SHEET = re.compile(r"[\[\]\*\?/\\:]")


def _num_format(ctype: str, values: Sequence[Any]) -> Optional[str]:
    if ctype == "integer":
        return "#,##0;[Red]-#,##0"
    if ctype == "percent":
        decs = max((-(v.as_tuple().exponent) for v in values if isinstance(v, Decimal)), default=2)
        return "0.0%" if decs <= 3 else "0.00%"
    if ctype == "number":
        decs = [max(0, -(v.as_tuple().exponent)) for v in values if isinstance(v, Decimal)]
        d = max(decs) if decs else 2
        d = min(d, 4)
        frac = ("." + "0" * d) if d else ""
        return f"#,##0{frac};[Red]-#,##0{frac}"
    if ctype == "date":
        return "yyyy-mm-dd"
    return None


def _cell_value(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    return v


def _excel_safe_text(v: Any) -> Any:
    """Identity for values; formula protection is applied by _guard after assignment."""
    return v


def _guard(cell) -> None:
    """Never let extracted text become a formula: openpyxl treats strings starting with '=' as formulas."""
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"


def _issue_map(issues: Sequence[ValidationIssue]) -> dict[tuple[str, int, Optional[int]], list[ValidationIssue]]:
    m: dict[tuple[str, int, Optional[int]], list[ValidationIssue]] = {}
    for i in issues:
        if i.table_id and i.row_index is not None and i.severity in ("critical", "warning"):
            m.setdefault((i.table_id, i.row_index, i.col_index), []).append(i)
    return m


def _unique_table_name(base: str, used: set[str]) -> str:
    name = re.sub(r"[^A-Za-z0-9_]", "_", base)
    if not name or not name[0].isalpha():
        name = "Tbl_" + name
    cand, k = name, 2
    while cand.lower() in used:
        cand = f"{name}_{k}"
        k += 1
    used.add(cand.lower())
    return cand


def _write_table(ws, t: LogicalTable, start_row: int, issues_by_cell, used_names: set[str], with_title: bool) -> int:
    """Write one table at start_row. Returns the next free row."""
    r0 = start_row
    if with_title:
        ws.cell(r0, 1, f"{t.table_id} — {t.title or 'Untitled table'}").font = TITLE_FONT
        ws.cell(r0 + 1, 1, f"Source pages: {', '.join(map(str, t.pages))}  ·  Method: {', '.join(t.methods)}  ·  "
                           f"{len(t.data_row_indices())} data rows").font = META_FONT
        r0 += 2
    headers = ["Row ID"] + list(t.columns) + ["Source Page"]
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(r0, ci, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    formats = [_num_format(ty, [row[i] for row in t.values]) for i, ty in enumerate(t.col_types)]
    widths = [len(h) for h in headers]
    for ri, (vals, meta) in enumerate(zip(t.values, t.row_meta)):
        excel_row = r0 + 1 + ri
        ws.cell(excel_row, 1, meta.row_id).font = Font(color="64748B", size=9)
        for ci, v in enumerate(vals):
            cell = ws.cell(excel_row, ci + 2, _excel_safe_text(_cell_value(v)))
            _guard(cell)
            if formats[ci] and isinstance(v, (Decimal, date)):
                cell.number_format = formats[ci]
            if isinstance(v, (Decimal, date)):
                cell.alignment = Alignment(horizontal="right")
            widths[ci + 1] = max(widths[ci + 1], len(str(t.rows[ri][ci])))
            found = issues_by_cell.get((t.table_id, ri, ci), [])
            if found:
                worst = "critical" if any(i.severity == "critical" for i in found) else "warning"
                cell.fill = ISSUE_FILL[worst]
                cell.comment = Comment("\n".join(f"[{i.severity.upper()}] {i.explanation}" for i in found)[:1500], "DocuPilot AI")
        ws.cell(excel_row, len(headers), meta.page).alignment = Alignment(horizontal="center")
        if meta.kind == "summary":
            for ci in range(1, len(headers) + 1):
                ws.cell(excel_row, ci).font = Font(bold=True)
        # Row-level findings without a specific column: highlight the Row ID cell.
        row_issues = issues_by_cell.get((t.table_id, ri, None), [])
        if row_issues:
            c = ws.cell(excel_row, 1)
            c.fill = ISSUE_FILL["critical" if any(i.severity == "critical" for i in row_issues) else "warning"]
            c.comment = Comment("\n".join(i.explanation for i in row_issues)[:1500], "DocuPilot AI")
    last_row = r0 + len(t.values)
    if t.values:
        ref = f"A{r0}:{get_column_letter(len(headers))}{last_row}"
        tab = Table(displayName=_unique_table_name(f"{t.table_id}_data", used_names), ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tab)
    for ci, w in enumerate(widths, start=1):
        cur = ws.column_dimensions[get_column_letter(ci)].width or 0
        ws.column_dimensions[get_column_letter(ci)].width = max(cur, min(max(w + 2, 8), 60))
    return last_row + 3


def _autosize(ws, max_width: int = 80) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), min(len(str(c.value)), max_width))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(10, w + 2)


def _header_row(ws, headers: Sequence[str]) -> None:
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(1, ci, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def build_workbook(result: ProcessingResult) -> bytes:
    """Create the .xlsx workbook and return its bytes."""
    wb = Workbook()
    issues_by_cell = _issue_map(result.issues)
    used_names: set[str] = set()

    # 1. Extracted Data
    ws = wb.active
    ws.title = "Extracted Data"
    if not result.tables:
        ws["A1"] = "No tables were detected in this document."
        ws["A1"].font = TITLE_FONT
    elif len(result.tables) == 1:
        _write_table(ws, result.tables[0], 1, issues_by_cell, used_names, with_title=False)
        ws.freeze_panes = "B2"
    else:
        row = 1
        for t in result.tables:
            row = _write_table(ws, t, row, issues_by_cell, used_names, with_title=True)
        ws.freeze_panes = "B4"

    # 2. Validation Report
    vr = wb.create_sheet("Validation Report")
    headers = ["Check", "Severity", "Status", "Confirmed", "Table", "Row", "Source Page", "Column",
               "Extracted Value", "Expected Value", "Difference", "Explanation", "Suggested Action"]
    _header_row(vr, headers)
    order = {"critical": 0, "warning": 1, "info": 2, "passed": 3}
    for i, iss in enumerate(sorted(result.issues, key=lambda x: (order.get(x.severity, 9), x.table_id or "", x.row_index or 0)), start=2):
        vals = [iss.check, iss.severity.title(), iss.status, "Yes" if iss.confirmed else ("—" if iss.severity == "passed" else "Suspected"),
                iss.table_id, iss.row_id, iss.page, iss.column, iss.extracted, iss.expected, iss.difference,
                iss.explanation, iss.suggested_action]
        for ci, v in enumerate(vals, start=1):
            c = vr.cell(i, ci, _excel_safe_text(v))
            _guard(c)
            if ci in (9, 10, 11) and isinstance(v, float):
                c.number_format = "#,##0.00;[Red]-#,##0.00"
        vr.cell(i, 2).fill = ISSUE_FILL.get(iss.severity, PatternFill())
    if result.issues:
        vr.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(result.issues) + 1}"
    _autosize(vr, max_width=70)

    # 3. Extraction Audit
    au = wb.create_sheet("Extraction Audit")
    headers = ["Source Page", "Table", "Row ID", "Row Type", "Extraction Method", "Page Fragment",
               "Normalisation / Correction Operations", "Warnings"]
    _header_row(au, headers)
    r = 2
    row_issue_text: dict[tuple[str, int], list[str]] = {}
    for iss in result.issues:
        if iss.table_id and iss.row_index is not None and iss.severity in ("critical", "warning"):
            row_issue_text.setdefault((iss.table_id, iss.row_index), []).append(f"{iss.check}: {iss.explanation}")
    for t in result.tables:
        for ri, m in enumerate(t.row_meta):
            warn = m.warnings + row_issue_text.get((t.table_id, ri), [])
            vals = [m.page, t.table_id, m.row_id, m.kind, m.method, m.segment, "; ".join(m.ops) or "none",
                    "; ".join(warn) or ""]
            for ci, v in enumerate(vals, start=1):
                _guard(au.cell(r, ci, _excel_safe_text(v)))
            r += 1
    if r > 2:
        au.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{r - 1}"
    _autosize(au, max_width=90)

    # 4. Document Summary
    ds = wb.create_sheet("Document Summary")
    ds["A1"] = "DocuPilot AI — extraction summary"
    ds["A1"].font = TITLE_FONT
    m = result.metrics()
    info = [("File", result.file_name), ("SHA-256", result.file_hash), ("Detected document type", result.doc_type),
            ("Pages", m["pages"]), ("Tables", m["tables"]), ("Data rows", m["rows"]), ("Pages using OCR", m["ocr_pages"]),
            ("Pages that failed", m["failed_pages"]), ("Processing time (s)", m["elapsed"]),
            ("Local LLM used", f"yes ({result.llm_model})" if result.llm_used else "no")]
    for i, (k, v) in enumerate(info, start=3):
        ds.cell(i, 1, k).font = Font(bold=True)
        ds.cell(i, 2, v)
    row = len(info) + 5
    ds.cell(row, 1, "Pages").font = TITLE_FONT
    row += 1
    for ci, h in enumerate(["Page", "Classification", "Method", "OCR", "OCR confidence", "Tables", "Status", "Notes"], start=1):
        c = ds.cell(row, ci, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for p in result.page_reports:
        row += 1
        for ci, v in enumerate([p.page, p.kind, p.method, "yes" if p.ocr_used else "no", p.ocr_confidence,
                                p.tables_found, p.status, "; ".join(p.warnings)], start=1):
            ds.cell(row, ci, v)
    row += 2
    ds.cell(row, 1, "Multi-page merge decisions").font = TITLE_FONT
    row += 1
    for ci, h in enumerate(["Table", "Earlier fragment", "Later fragment", "Merged", "Score", "Evidence", "Uncertainty"], start=1):
        c = ds.cell(row, ci, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    if not result.merge_decisions:
        row += 1
        ds.cell(row, 1, "No cross-page table boundaries were evaluated.")
    for d in result.merge_decisions:
        row += 1
        for ci, v in enumerate([d.table_id, d.left, d.right, "yes" if d.merged else "no", d.score,
                                "; ".join(d.evidence), d.uncertainty or ""], start=1):
            ds.cell(row, ci, v)
    row += 2
    ds.cell(row, 1, "Header normalisation").font = TITLE_FONT
    row += 1
    for ci, h in enumerate(["Table", "Original header", "Normalised header", "Semantic role", "Column type"], start=1):
        c = ds.cell(row, ci, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for t in result.tables:
        for orig, new, role, ty in zip(t.original_header, t.columns, t.roles, t.col_types):
            row += 1
            for ci, v in enumerate([t.table_id, orig, new, role or "", ty], start=1):
                _guard(ds.cell(row, ci, _excel_safe_text(v)))
    row += 2
    ds.cell(row, 1, "Fields found outside tables").font = TITLE_FONT
    row += 1
    for ci, h in enumerate(["Field", "Label in document", "Value", "Page", "Source"], start=1):
        c = ds.cell(row, ci, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for f in result.fields:
        row += 1
        for ci, v in enumerate([f.key, f.label, _cell_value(f.value), f.page, f.source], start=1):
            _guard(ds.cell(row, ci, _excel_safe_text(v)))
    _autosize(ds, max_width=100)

    # 5+. One sheet per table when there are several.
    if len(result.tables) > 1:
        for t in result.tables:
            name = _SAFE_SHEET.sub("_", f"{t.table_id} {t.title or ''}".strip())[:31]
            tw = wb.create_sheet(name)
            _write_table(tw, t, 1, issues_by_cell, used_names, with_title=False)
            tw.freeze_panes = "B2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
